"""
ingest/adapters/remote_table.py - ONE adapter for every "table somewhere on the web" source
==========================================================================================

USED FOR (registry kind: remote_table)
    Every dataset onboarded by the onboarding agent (onboard/): a CSV / Excel / JSON file (for example a CKAN
    resource on Bayanat), a JSON API endpoint, or an HTML <table> on a page. The PATTERN is generic; the
    exact location is fixed in the registry entry, so a run never guesses:

        request:
          url:         where to GET (or POST) the data
          format:      csv | xlsx | json | html
          json_path:   for json: dotted path to the list of records ("result.records", "" = the top level)
          table_index: for html: which <table> on the page (0-based), checked against `columns`
          sheet:       for xlsx: sheet name (default: the first sheet)
          method / body / headers:   for a JSON call captured from a page (POST bodies are replayed as is)
          page_param / max_pages:    optional simple paging: ?<page_param>=1,2,3... until an empty page
          field_map:   which column is the name / date / status / category / website / key (used by the
                       `generic` normaliser in ingest/schemas.py)

    parse() is PURE (bytes in, rows out): the onboarding agent uses the same function for its trial fetch,
    so what the person approves is exactly what the scheduled runs will read.

ERRORS
    A file that no longer parses, a json_path that no longer leads to a list, or a table whose header changed
    raises schema_drift: the old snapshot keeps serving and the run log says what changed.
"""
from __future__ import annotations

import csv
import io
import json
from urllib.parse import urlencode, urlparse, parse_qsl, urlunparse

from ..responses import FetchError
from .base import Context, FetchResult
from .html_table import _Tables

MAX_PAGES = 50


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    return " ".join(str(v).split())


def _rows_from_matrix(matrix: list[list], expect: list[str] | None = None) -> list[dict]:
    """First row with 2+ non-empty cells is the header; empty rows are dropped; duplicate headers get a suffix."""
    start = next((i for i, r in enumerate(matrix) if sum(1 for c in r if _cell(c)) >= 2), None)
    if start is None:
        return []
    header, seen = [], {}
    for h in matrix[start]:
        h = _cell(h) or "column"
        seen[h] = seen.get(h, 0) + 1
        header.append(h if seen[h] == 1 else f"{h} ({seen[h]})")
    if expect and not all(e in header for e in expect):
        raise FetchError("schema_drift", f"expected columns {expect}, found {header[:12]}")
    out = []
    for r in matrix[start + 1:]:
        vals = [_cell(c) for c in r]
        if any(vals):
            out.append({header[i]: (vals[i] if i < len(vals) else "") for i in range(len(header))})
    return out


def walk(obj, path: str):
    """Follow a dotted path ('result.records', 'data.items.0'); '' returns obj itself."""
    for part in [p for p in (path or "").split(".") if p]:
        if isinstance(obj, list) and part.isdigit():
            obj = obj[int(part)]
        elif isinstance(obj, dict) and part in obj:
            obj = obj[part]
        else:
            raise FetchError("schema_drift", f"json path '{path}' broke at '{part}'")
    return obj


def flatten(d: dict, prefix: str = "") -> dict:
    """{'a': {'b': 1}, 'c': [1, 2]} -> {'a.b': '1', 'c': '1; 2'}: every record becomes one flat row of text."""
    out = {}
    for k, v in d.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(flatten(v, key + "."))
        elif isinstance(v, list):
            out[key] = "; ".join(_cell(x) if not isinstance(x, (dict, list)) else json.dumps(x, ensure_ascii=False)
                                 for x in v)
        else:
            out[key] = _cell(v)
    return out


def parse(body: bytes, fmt: str, req: dict) -> list[dict]:
    """Bytes -> list of flat rows (all values text). Raises schema_drift if the content is not what req says."""
    try:
        if fmt == "csv":
            text = body.decode("utf-8-sig", errors="replace")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            return _rows_from_matrix(list(csv.reader(io.StringIO(text), dialect)), req.get("columns"))
        if fmt == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            ws = wb[req["sheet"]] if req.get("sheet") else wb.worksheets[0]
            return _rows_from_matrix([list(r) for r in ws.iter_rows(values_only=True)], req.get("columns"))
        if fmt == "json":
            data = walk(json.loads(body.decode("utf-8-sig")), req.get("json_path", ""))
            if not isinstance(data, list):
                raise FetchError("schema_drift", f"json path '{req.get('json_path', '')}' is not a list")
            return [flatten(x) for x in data if isinstance(x, dict)]
        if fmt == "html":
            p = _Tables()
            p.feed(body.decode("utf-8", errors="replace"))
            i = int(req.get("table_index", 0))
            if i >= len(p.tables):
                raise FetchError("schema_drift", f"table {i} not found ({len(p.tables)} tables on the page)")
            return _rows_from_matrix(p.tables[i], req.get("columns"))
    except FetchError:
        raise
    except Exception as e:                                   # corrupt file, wrong format, bad JSON
        raise FetchError("schema_drift", f"could not read the {fmt} content: {type(e).__name__}: {e}") from e
    raise FetchError("config", f"unknown format {fmt!r}")


def _with_page(url: str, param: str, n: int) -> str:
    u = urlparse(url)
    q = dict(parse_qsl(u.query))
    q[param] = str(n)
    return urlunparse(u._replace(query=urlencode(q)))


def fetch(spec, ctx: Context) -> FetchResult:
    req = spec.request
    http = ctx.session.http                                  # the run's HTTP client (logging, retries, budget)
    pages = [1] if not req.get("page_param") else range(1, int(req.get("max_pages", MAX_PAGES)) + 1)
    rows, last_first = [], None
    for n in pages:
        url = _with_page(req["url"], req["page_param"], n) if req.get("page_param") else req["url"]
        resp = http.request(spec.id, req.get("method", "GET"), url, headers=req.get("headers"),
                            json_body=req.get("body"), purpose=f"{req['format']} page {n}")
        ctx.raw.save(resp, "page")
        if resp.status != 200:
            raise FetchError("http_status", f"HTTP {resp.status} for {url}", status=resp.status)
        got = parse(resp.body, req["format"], req)
        ctx.runlog.event(spec.id, "page", "info", f"page {n}: {len(got)} rows", bytes=len(resp.body))
        if not got or (got and got[0] == last_first):            # empty or repeated page = the end
            break
        last_first = got[0]
        rows += got
    return FetchResult(items=rows, pages=len(rows) and n, source_url=spec.source_page, total_reported=None)
